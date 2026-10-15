from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/rahulsolanki/Documents/climate-resilient-workforce-agent')
OUT = ROOT / '.docx-output' / 'Baseline-Planning-Cycle-Schema-Guide-Final.docx'
OUT.parent.mkdir(exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].font.color.rgb = RGBColor(35, 45, 42)
styles['Normal'].paragraph_format.space_after = Pt(5)
for style_name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 11)]:
    st = styles[style_name]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(8 if style_name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)
    if style_name == 'Title' and st._element.pPr is not None:
        borders = st._element.pPr.find(qn('w:pBdr'))
        if borders is not None:
            st._element.pPr.remove(borders)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color='23302C', size=8.3):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(headers, rows, widths=None, font_size=8.1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF', font_size)
        shade(t.rows[0].cells[i], '245C4B')
    header_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    header_pr.append(repeat)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, i == 0, size=font_size)
            if row_index % 2 == 1:
                shade(cells[i], 'F0F4F1')
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def code_block(text, size=7.2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F2F5F3')
    ppr.append(shd)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(30, 48, 41)


def para(text='', bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


# Page 1: mental model and core contract
p = doc.add_paragraph(style='Title')
p.add_run('Baseline Planning Cycle Schema Guide')
direct_borders = p._p.get_or_add_pPr().find(qn('w:pBdr'))
if direct_borders is not None:
    p._p.get_or_add_pPr().remove(direct_borders)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(11)
r = subtitle.add_run('A short field guide with a complete example')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(82, 99, 91)

para('A baseline is the original workforce schedule saved before climate-aware planning changes it. A planning cycle stores that schedule together with the exact work orders, workers, and equipment supplied for that planning period. This gives later schedule evaluation a stable reference.')

doc.add_heading('How the records fit together', level=1)
table(
    ['Record', 'What it represents', 'In the example'],
    [
        ('PlanningCycleRequest', 'The create request: one cycle ID, its input dataset, original assignments, and any unscheduled work.', 'Input sent to POST /api/planning-cycles'),
        ('PlanningDataset', 'The source work orders, workers, and equipment captured for this planning cycle.', '1 work order, 2 workers, 1 equipment item'),
        ('BaselineSnapshot', 'The persisted, write-once copy of the original schedule and its evaluation.', 'Returned by the create endpoint and available from the baseline GET endpoint'),
        ('ScheduleMetrics', 'Versioned counts and time totals calculated by the shared evaluator.', '1 assignment, 2 worker allocations, 180 scheduled minutes'),
    ], [1.35, 3.1, 2.55], font_size=8.2)

doc.add_heading('Reading the lifecycle', level=1)
para('The client submits the original schedule before optimization. The service checks that each assigned work order, worker, and equipment ID exists in the captured dataset, calculates metrics, and saves one baseline per planning_cycle_id. A duplicate cycle ID is rejected, so the original reference is not overwritten.')
para('The example uses Asia/Kolkata timestamps. A TimeWindow includes its start and excludes its end. The scheduled-minute metric is the difference between the assignment window end and start; it does not subtract cooling breaks.')

# Page 2: complete input example
doc.add_page_break()
doc.add_heading('Example create request', level=1)
para('This compact example schedules one three-hour maintenance task with a two-person crew and one crane. All IDs referenced by the assignment are present in the dataset.')
example = '''{
  "planning_cycle_id": "CYCLE-2048",
  "dataset": {
    "work_orders": [{
      "work_order_id": "WO-104",
      "title": "Replace heat-exchanger gasket",
      "duration_minutes": 180,
      "priority": "high",
      "site_id": "VAD-PLANT-01",
      "required_skills": ["mechanical"],
      "required_certifications": [],
      "min_crew_size": 2,
      "max_crew_size": 2,
      "depends_on": [],
      "required_equipment_ids": ["EQ-CRANE-01"],
      "planning_window": {"start": "2026-10-13T06:00:00+05:30", "end": "2026-10-13T14:00:00+05:30"},
      "deadline": "2026-10-13T14:00:00+05:30",
      "exposure_type": "outdoor_heat",
      "expected_exposure_duration_minutes": 150
    }],
    "workers": [
      {"worker_id": "W-01", "name": "Asha Patel", "skills": ["mechanical"],
       "certifications": [],
       "availability_windows": [{"start": "2026-10-13T06:00:00+05:30", "end": "2026-10-13T14:00:00+05:30"}],
       "break_requirements": {"min_break_minutes": 15, "max_continuous_work_minutes": 120, "cooling_break_minutes": 20},
       "max_shift_duration_minutes": 480},
      {"worker_id": "W-02", "name": "Ravi Shah", "skills": ["mechanical"],
       "certifications": [],
       "availability_windows": [{"start": "2026-10-13T06:00:00+05:30", "end": "2026-10-13T14:00:00+05:30"}],
       "break_requirements": {"min_break_minutes": 15, "max_continuous_work_minutes": 120, "cooling_break_minutes": 20},
       "max_shift_duration_minutes": 480}
    ],
    "equipment": [{
      "equipment_id": "EQ-CRANE-01",
      "equipment_type": "mobile-crane",
      "availability_windows": [{"start": "2026-10-13T06:00:00+05:30", "end": "2026-10-13T14:00:00+05:30"}],
      "operational_status": "available",
      "compatible_task_types": ["maintenance"],
      "inspection_valid_until": "2026-12-31T23:59:00+05:30",
      "site_id": "VAD-PLANT-01"
    }]
  },
  "assignments": [{
    "assignment_id": "A-104",
    "work_order_id": "WO-104",
    "site_id": "VAD-PLANT-01",
    "worker_ids": ["W-01", "W-02"],
    "equipment_ids": ["EQ-CRANE-01"],
    "scheduled_window": {"start": "2026-10-13T06:00:00+05:30", "end": "2026-10-13T09:00:00+05:30"},
    "cooling_breaks": [],
    "decision": "proceed",
    "rationale": "Original dispatch schedule before climate-aware planning."
  }],
  "unscheduled": []
}'''
code_block(example, size=7.0)
para('The request leaves out optional fields only where the schemas define a default or allow null. For example, WorkOrder.is_weather_exposed is inferred from exposure_type when omitted. The example does not include that optional field.')

# Page 3: dataset field reference
doc.add_page_break()
doc.add_heading('Dataset fields', level=1)
para('PlanningDataset groups three lists. The list items are the existing WorkOrder, Worker, and Equipment models; they describe the resources and constraints available to the planning cycle.')

doc.add_heading('WorkOrder', level=2)
table(['Field', 'Meaning'], [
    ('work_order_id, title', 'Stable task identifier and readable task name.'),
    ('duration_minutes, priority', 'Expected task duration and urgency: critical, high, medium, or low.'),
    ('site_id', 'Site where the task takes place.'),
    ('required_skills, required_certifications', 'Qualifications needed to perform the task.'),
    ('min_crew_size, max_crew_size', 'Allowed crew-size range; maximum cannot be below minimum.'),
    ('depends_on', 'Work-order IDs that must precede this task.'),
    ('required_equipment_ids', 'Equipment IDs required by the task.'),
    ('planning_window', 'Timezone-aware earliest/latest allowed interval for planning.'),
    ('deadline', 'Optional timezone-aware due time.'),
    ('exposure_type', 'Work hazard category, such as outdoor_heat or confined_space.'),
    ('expected_exposure_duration_minutes', 'Expected exposure time; cannot exceed task duration.'),
    ('exposure_windows', 'Optional intervals when exposure is expected.'),
    ('is_weather_exposed', 'Optional explicit weather-exposure flag. Inferred from exposure_type when omitted.'),
], [2.0, 5.0], font_size=7.8)

doc.add_heading('Worker', level=2)
table(['Field', 'Meaning'], [
    ('worker_id, name', 'Stable worker identifier and display name.'),
    ('skills, certifications', 'Worker qualifications; certification may include a timezone-aware expiry.'),
    ('availability_windows', 'One or more timezone-aware periods when the worker is available.'),
    ('break_requirements', 'Minimum break length, maximum continuous work time, and cooling-break length, all in minutes.'),
    ('max_shift_duration_minutes', 'Maximum allowed shift length in minutes.'),
], [2.0, 5.0], font_size=7.8)

doc.add_heading('Equipment', level=2)
table(['Field', 'Meaning'], [
    ('equipment_id, equipment_type', 'Stable equipment identifier and equipment category.'),
    ('availability_windows', 'Timezone-aware periods when the equipment is available.'),
    ('operational_status', 'Current state, such as available, in_use, maintenance, or out_of_service.'),
    ('compatible_task_types', 'Task categories this equipment can support.'),
    ('inspection_valid_until', 'Optional timezone-aware inspection expiry.'),
    ('site_id', 'Optional home or assigned site for the equipment.'),
], [2.0, 5.0], font_size=7.8)

# Page 4: schedule, snapshot, metrics
doc.add_page_break()
doc.add_heading('Schedule and snapshot fields', level=1)
doc.add_heading('ScheduleAssignment', level=2)
table(['Field', 'Meaning'], [
    ('assignment_id', 'Unique identifier for this scheduled task instance.'),
    ('work_order_id, site_id', 'Task and site this assignment belongs to.'),
    ('worker_ids, equipment_ids', 'Worker and equipment identifiers allocated to the task.'),
    ('scheduled_window', 'Timezone-aware start and end of the scheduled interval.'),
    ('cooling_breaks', 'Break intervals and the reason for each break. Empty list means none are recorded.'),
    ('decision', 'Schedule decision: proceed, reschedule, extra_breaks, or blocked.'),
    ('rationale', 'Short explanation for the decision in this source schedule.'),
], [2.0, 5.0], font_size=8)

doc.add_heading('Other records', level=2)
table(['Schema / field', 'Meaning'], [
    ('UnscheduledWorkOrder.work_order_id', 'Work order that has no assignment.'),
    ('UnscheduledWorkOrder.reasons', 'One or more reasons it was not scheduled.'),
    ('BaselineSnapshot.snapshot_id', 'Unique identifier for the captured snapshot.'),
    ('BaselineSnapshot.planning_cycle_id', 'Cycle ID that owns the snapshot.'),
    ('BaselineSnapshot.captured_at', 'Timezone-aware time the baseline was stored.'),
    ('BaselineSnapshot.dataset', 'Copy of the exact input dataset used for evaluation.'),
    ('BaselineSnapshot.assignments / unscheduled', 'Original schedule records, including unscheduled work orders.'),
    ('BaselineSnapshot.metrics', 'Metric values calculated for the saved baseline.'),
], [2.5, 4.5], font_size=8)

doc.add_heading('Sample metrics from the example', level=2)
table(['Metric', 'Value', 'How it is calculated'], [
    ('definition_version', '1', 'Identifies the evaluator’s metric definition version.'),
    ('input_work_order_count', '1', 'Number of work orders in the captured dataset.'),
    ('input_worker_count', '2', 'Number of workers in the captured dataset.'),
    ('input_equipment_count', '1', 'Number of equipment items in the captured dataset.'),
    ('assignment_count', '1', 'Number of assignment records.'),
    ('scheduled_work_order_count', '1', 'Distinct work orders referenced by assignments.'),
    ('worker_allocation_count', '2', 'Sum of worker IDs across assignments.'),
    ('equipment_allocation_count', '1', 'Sum of equipment IDs across assignments.'),
    ('scheduled_minutes', '180', '09:00 minus 06:00; breaks are not subtracted.'),
    ('unscheduled_work_order_count', '0', 'Number of entries in the unscheduled list.'),
], [2.05, 0.7, 4.25], font_size=7.6)

doc.add_heading('What the metrics do not say', level=2)
para('These version 1 metrics summarize counts and scheduled time. They do not measure safety compliance, schedule quality, improvements, or recommendations. The evaluator checks that IDs referenced by assignments exist in the captured dataset; it does not replace full scheduling-constraint validation.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = footer.add_run('Baseline Planning Cycle Schema Guide')
fr.font.name = 'Arial'
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(105, 118, 111)

doc.save(OUT)
print(OUT)
